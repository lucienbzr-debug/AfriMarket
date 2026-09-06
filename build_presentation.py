"""Génère la présentation PowerPoint de direction AfriMarket (charte graphique officielle)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BASE = r"c:\Users\Lbzr\Desktop\FORMATION CLAUDE & AI BY AKSANTE\Jour 3\AfriMarket"
FIG = fr"{BASE}\rapport\figures"
LOGO = fr"{BASE}\Logo AfriMarket.png"
OUT = fr"{BASE}\rapport\AfriMarket_Presentation_Direction.pptx"

# ---------- Charte graphique (couleurs officielles extraites du logo) ----------
RED = RGBColor(0xDC, 0x3D, 0x2A)
DARK = RGBColor(0x2A, 0x2A, 0x28)
GRAY = RGBColor(0x9A, 0x9A, 0x97)
LIGHT_GRAY = RGBColor(0xF3, 0xF2, 0xF1)
GOLD = RGBColor(0xE8, 0xA3, 0x3D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def fill_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line:
        shp.line.color.rgb = color
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size=18, color=DARK, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
    return box


def add_bullets(slide, x, y, w, h, items, size=13.5, color=DARK, bullet_color=RED, space_after=10, bold_first=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = 1.08
        r1 = p.add_run()
        r1.text = "●  "
        r1.font.size = Pt(size)
        r1.font.color.rgb = bullet_color
        r1.font.bold = True
        r1.font.name = FONT
        if isinstance(item, tuple):
            title, rest = item
            r2 = p.add_run()
            r2.text = title
            r2.font.size = Pt(size)
            r2.font.color.rgb = RED
            r2.font.bold = True
            r2.font.name = FONT
            r3 = p.add_run()
            r3.text = rest
            r3.font.size = Pt(size)
            r3.font.color.rgb = color
            r3.font.name = FONT
        else:
            r2 = p.add_run()
            r2.text = item
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.name = FONT
    return box


def add_logo(slide, x=0.45, y=0.35, h=0.55):
    slide.shapes.add_picture(LOGO, Inches(x), Inches(y), height=Inches(h))


def add_footer(slide, page_no):
    add_rect(slide, 0, 7.12, 13.333, 0.38, DARK)
    add_text(slide, 0.45, 7.14, 8, 0.35, "AfriMarket — Résumé Exécutif à destination de la Direction",
              size=9.5, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 12.0, 7.14, 0.9, 0.35, str(page_no), size=9.5, color=GRAY,
              align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def add_header(slide, title, kicker=None):
    add_rect(slide, 0, 0, 13.333, 1.15, WHITE)
    add_rect(slide, 0, 1.15, 13.333, 0.05, RED)
    add_logo(slide, x=11.9, y=0.32, h=0.5)
    if kicker:
        add_text(slide, 0.6, 0.18, 9, 0.3, kicker.upper(), size=11.5, color=RED, bold=True)
        add_text(slide, 0.6, 0.48, 9, 0.6, title, size=24, color=DARK, bold=True)
    else:
        add_text(slide, 0.6, 0.33, 9, 0.6, title, size=26, color=DARK, bold=True)


def add_picture_fit(slide, path, x, y, w, h):
    from PIL import Image
    im = Image.open(path)
    iw, ih = im.size
    target_ratio = w / h
    src_ratio = iw / ih
    if src_ratio > target_ratio:
        new_w, new_h = w, w / src_ratio
    else:
        new_h, new_w = h, h * src_ratio
    px = x + (w - new_w) / 2
    py = y + (h - new_h) / 2
    slide.shapes.add_picture(path, Inches(px), Inches(py), width=Inches(new_w), height=Inches(new_h))


# =========================================================================
# SLIDE 1 — COUVERTURE
# =========================================================================
s = add_slide()
fill_bg(s, DARK)
add_rect(s, 0, 6.55, 13.333, 0.12, RED)
slide_logo = s.shapes.add_picture(LOGO, Inches(5.42), Inches(1.15), height=Inches(1.9))
add_text(s, 1, 3.35, 11.33, 0.9, "Dashboard Stratégique & Analyse Commerciale",
         size=32, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, 1, 4.15, 11.33, 0.6, "Résumé exécutif à destination de la Direction",
         size=17, color=GOLD, align=PP_ALIGN.CENTER, italic=True)
add_text(s, 1, 5.9, 11.33, 0.4, "Période analysée : Juillet — Décembre 2025",
         size=13, color=GRAY, align=PP_ALIGN.CENTER)
add_text(s, 1, 6.2, 11.33, 0.4, "9 400 commandes nettoyées · 8 villes · 4 catégories · 1 747 clients",
         size=12, color=GRAY, align=PP_ALIGN.CENTER)

# =========================================================================
# SLIDE 2 — CONTEXTE & OBJECTIFS
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "Contexte & objectifs de l'analyse", kicker="Introduction")
add_text(s, 0.6, 1.55, 12.1, 0.9,
         "AfriMarket est une entreprise e-commerce panafricaine opérant dans 8 villes d'Afrique "
         "francophone, sur 4 catégories : Électronique, Mode, Beauté et Maison.",
         size=15, color=DARK, line_spacing=1.15)
add_text(s, 0.6, 2.55, 5.6, 0.4, "4 signaux remontés par la Direction", size=15, color=RED, bold=True)
add_bullets(s, 0.6, 3.0, 5.7, 3.6, [
    "Variations importantes du chiffre d'affaires selon les mois",
    "Taux de retour préoccupant sur certains produits",
    "Dépenses marketing élevées, ROI incertain par canal",
    "Écarts de performance importants selon les villes",
], size=14)
add_rect(s, 6.7, 2.55, 6.0, 3.9, LIGHT_GRAY)
add_text(s, 7.0, 2.75, 5.4, 0.4, "Portée de l'étude", size=15, color=RED, bold=True)
add_bullets(s, 7.0, 3.2, 5.4, 3.1, [
    "6 mois d'activité commerciale complète",
    "10 100 commandes brutes → 9 400 après nettoyage",
    "Analyse par catégorie, ville, canal marketing et client",
    "5 recommandations chiffrées et un plan d'action à 30 jours",
], size=14)
add_footer(s, 2)

# =========================================================================
# SLIDE 3 — KPIs CLES
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "Chiffres clés de la période", kicker="Vue d'ensemble")
kpis = [
    ("2,51 M$", "Chiffre d'affaires total"),
    ("430 k$", "Profit net estimé"),
    ("272,03 $", "Panier moyen"),
    ("1,9 %", "Taux d'annulation"),
    ("8,1 %", "Taux de retour"),
]
card_w, gap = 2.28, 0.2
start_x = (13.333 - (card_w * 5 + gap * 4)) / 2
for i, (val, label) in enumerate(kpis):
    x = start_x + i * (card_w + gap)
    is_accent = i == 0
    add_rect(s, x, 2.1, card_w, 2.3, RED if is_accent else LIGHT_GRAY)
    add_text(s, x, 2.5, card_w, 0.8, val, size=27, bold=True,
              color=WHITE if is_accent else DARK, align=PP_ALIGN.CENTER)
    add_text(s, x + 0.12, 3.35, card_w - 0.24, 0.8, label, size=12.5,
              color=WHITE if is_accent else GRAY, align=PP_ALIGN.CENTER)
add_picture_fit(s, fr"{FIG}\01_ca_mensuel.png", 1.0, 4.65, 11.3, 2.15)
add_footer(s, 3)

# =========================================================================
# SLIDE 4 — CATEGORIE
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "Quelle catégorie prioriser ?", kicker="Analyse produit")
add_picture_fit(s, fr"{FIG}\02_categorie.png", 0.5, 1.45, 12.3, 4.05)
add_rect(s, 0.5, 5.65, 12.3, 1.15, LIGHT_GRAY)
add_text(s, 0.75, 5.78, 11.8, 0.95,
         "Électronique génère 74,6 % du CA (1,87 M$) mais affiche la marge la plus faible (20 %) et le "
         "taux de retour le plus élevé (13,8 %). Beauté est la catégorie la plus rentable (marge 50 %) et "
         "la plus fiable (2,8 % de retour) mais reste marginale (3 % du CA).",
         size=13, color=DARK, line_spacing=1.1)
add_footer(s, 4)

# =========================================================================
# SLIDE 5 — GEOGRAPHIE
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "Où investir davantage ?", kicker="Analyse géographique")
add_picture_fit(s, fr"{FIG}\03_ville.png", 0.5, 1.45, 7.1, 4.9)
add_text(s, 7.9, 1.7, 4.9, 0.4, "Constats clés", size=15, color=RED, bold=True)
add_bullets(s, 7.9, 2.15, 4.9, 4.2, [
    ("Kinshasa — ", "domine largement (752 k$ de CA, 130 k$ de profit) : priorité d'investissement."),
    ("Douala — ", "meilleure croissance du réseau (+76 %) mais taux d'annulation anormal (12,9 % vs ~0 % ailleurs)."),
    ("Libreville — ", "recule de 46 % : à réévaluer avant tout nouveau budget marketing."),
    ("Brazzaville — ", "marché émergent en croissance (+60 %) à surveiller."),
], size=13.5)
add_footer(s, 5)

# =========================================================================
# SLIDE 6 — MARKETING
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "Quel canal mérite plus de budget ?", kicker="Analyse marketing")
add_picture_fit(s, fr"{FIG}\04_roi_canal.png", 0.5, 1.45, 12.3, 4.05)
add_rect(s, 0.5, 5.65, 12.3, 1.15, LIGHT_GRAY)
add_text(s, 0.75, 5.78, 11.8, 0.95,
         "Email affiche un ROI de 225x pour seulement 2 349 $ investis (le plus petit budget). Instagram "
         "Ads consomme le plus gros budget (37 637 $) pour un ROI de 24x seulement. Un rééquilibrage "
         "progressif vers Email et Google Ads (ROI 49x) maximiserait le retour marginal du budget.",
         size=13, color=DARK, line_spacing=1.1)
add_footer(s, 6)

# =========================================================================
# SLIDE 7 — CLIENTS
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "Comment améliorer la rétention ?", kicker="Analyse clients")
add_picture_fit(s, fr"{FIG}\05_pareto.png", 0.5, 1.45, 7.3, 4.9)
add_text(s, 8.0, 1.7, 4.85, 0.4, "Constats clés", size=15, color=RED, bold=True)
add_bullets(s, 8.0, 2.15, 4.85, 4.2, [
    ("1 747 clients, ", "dont 74 % déjà récurrents (plus d'une commande) : un socle de fidélité solide."),
    ("Effet Pareto marqué : ", "les 350 clients les plus rentables (20 % de la base) génèrent 64,5 % du CA total."),
    ("Implication : ", "la valeur de l'entreprise repose sur une minorité de clients à protéger en priorité "
     "via un programme de fidélisation dédié."),
], size=13.5)
add_footer(s, 7)

# =========================================================================
# SLIDE 8 — RECOMMANDATIONS
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "5 recommandations stratégiques", kicker="Plan d'action")
recos = [
    ("1. Sécuriser la marge Électronique", "Négocier les coûts fournisseurs et traiter la cause des retours : "
     "levier le plus élevé sur le profit total (74,6 % du CA, marge 20 %)."),
    ("2. Réduire le taux de retour Électronique", "Audit qualité/SAV ciblé — cause principale du taux de "
     "retour global de 8,1 %."),
    ("3. Réallouer le budget marketing", "Transférer 15-20 % du budget Instagram/Influenceur vers Email "
     "(ROI 225x) et Google Ads (ROI 49x), en test A/B avant généralisation."),
    ("4. Corriger l'anomalie de Douala", "Investiguer en urgence le taux d'annulation de 12,9 % (paiement, "
     "stock, livraison) ; geler tout budget supplémentaire à Libreville (-46 %)."),
    ("5. Lancer un programme de fidélisation VIP", "Protéger les 350 clients Pareto (64,5 % du CA) : le ROI "
     "de rétention y est supérieur à celui de l'acquisition."),
]
y = 1.55
for title, detail in recos:
    add_rect(s, 0.6, y, 0.09, 0.95, RED)
    add_text(s, 0.85, y, 11.9, 0.4, title, size=15, color=DARK, bold=True)
    add_text(s, 0.85, y + 0.42, 11.9, 0.55, detail, size=12, color=GRAY, line_spacing=1.05)
    y += 1.08
add_footer(s, 8)

# =========================================================================
# SLIDE 9 — PLAN D'ACTION 30 JOURS
# =========================================================================
s = add_slide()
fill_bg(s, WHITE)
add_header(s, "Plan d'action à 30 jours", kicker="Mise en œuvre")
steps = [
    ("Semaine 1", "Audit qualité/SAV Électronique + audit opérationnel des annulations à Douala."),
    ("Semaine 2", "Transfert de 15-20 % du budget Instagram Ads/Influenceur vers Email/Google Ads, "
     "en test A/B sur un mois."),
    ("Semaines 3-4", "Lancement pilote du programme de fidélisation VIP sur les 350 clients Pareto identifiés."),
]
card_w = 3.85
gap = 0.35
start_x = (13.333 - (card_w * 3 + gap * 2)) / 2
for i, (label, detail) in enumerate(steps):
    x = start_x + i * (card_w + gap)
    add_rect(s, x, 1.8, card_w, 0.75, RED)
    add_text(s, x, 1.8, card_w, 0.75, label, size=17, color=WHITE, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, x, 2.55, card_w, 2.9, LIGHT_GRAY)
    add_text(s, x + 0.25, 2.8, card_w - 0.5, 2.5, detail, size=13, color=DARK, line_spacing=1.15)
add_rect(s, 1.0, 5.85, 11.33, 1.0, DARK)
add_text(s, 1.3, 6.02, 10.8, 0.7,
         "Ces trois actions ciblent directement les 4 signaux remontés par la Direction, sans budget "
         "supplémentaire net : une réallocation de ressources existantes.",
         size=13, color=WHITE, italic=True, anchor=MSO_ANCHOR.MIDDLE)
add_footer(s, 9)

# =========================================================================
# SLIDE 10 — CONCLUSION / MERCI
# =========================================================================
s = add_slide()
fill_bg(s, DARK)
add_rect(s, 0, 0, 13.333, 0.12, RED)
s.shapes.add_picture(LOGO, Inches(5.92), Inches(0.9), height=Inches(1.3))
add_text(s, 1, 2.5, 11.33, 1.0,
         "Un modèle rentable, trois leviers de rentabilité identifiés",
         size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, 1.5, 3.55, 10.33, 1.3,
         "2,51 M$ de CA et 430 k$ de profit net estimé sur 6 mois — avec des fuites de rentabilité "
         "claires et quantifiées sur les retours Électronique, l'allocation marketing et l'anomalie "
         "opérationnelle de Douala.",
         size=14, color=GRAY, align=PP_ALIGN.CENTER, line_spacing=1.2)
add_text(s, 1, 6.3, 11.33, 0.5, "Merci de votre attention", size=18, color=GOLD, bold=True,
          align=PP_ALIGN.CENTER, italic=True)

prs.save(OUT)
print("Présentation sauvegardée :", OUT)
