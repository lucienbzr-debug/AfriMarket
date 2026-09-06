from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"c:\Users\Lbzr\Desktop\FORMATION CLAUDE & AI BY AKSANTE\Jour 3\AfriMarket"
FIG = fr"{BASE}\rapport\figures"
OUT = fr"{BASE}\rapport\Resume_Executif_AfriMarket.docx"

NAVY = RGBColor(0x0D, 0x1B, 0x2A)
BLUE = RGBColor(0x15, 0x65, 0xC0)
GRAY = RGBColor(0x45, 0x45, 0x45)

doc = Document()

# base font
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.font.color.rgb = GRAY

for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def shade_cell(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_title(text, size=20, color=NAVY, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return p


def add_h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = BLUE
    return p


def add_body(text, bold=False, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = GRAY
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = GRAY
    return p


# ---------- PAGE 1 : COUVERTURE / CONTEXTE ----------
add_title("AfriMarket — Résumé Exécutif", size=22)
p = doc.add_paragraph()
r = p.add_run("Analyse stratégique des données commerciales — Juillet à Décembre 2025")
r.italic = True
r.font.size = Pt(12)
r.font.color.rgb = BLUE
p.paragraph_format.space_after = Pt(14)

add_h2("Contexte")
add_body(
    "AfriMarket est une entreprise e-commerce panafricaine opérant dans 8 villes d'Afrique francophone, "
    "vendant des produits dans 4 catégories : Électronique, Mode, Beauté et Maison. La direction a constaté "
    "des variations importantes du chiffre d'affaires, un taux de retour préoccupant sur certains produits, "
    "des dépenses marketing élevées, et des différences de performance selon les villes. Cette analyse "
    "couvre 6 mois d'activité commerciale (10 100 commandes brutes, réduites à 9 400 commandes après "
    "nettoyage) et répond à ces quatre signaux avec des recommandations chiffrées."
)

add_h2("Chiffres clés (KPIs globaux)")

kpis = [
    ("CA total", "2 507 325 $"),
    ("Profit net estimé", "430 368 $"),
    ("Panier moyen", "272,03 $"),
    ("Taux d'annulation", "1,9 %"),
    ("Taux de retour", "8,1 %"),
]
table = doc.add_table(rows=1, cols=len(kpis))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
for i, (label, val) in enumerate(kpis):
    cell = hdr[i]
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(val)
    r1.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = BLUE
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(label)
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GRAY
    shade_cell(cell, "F0F4F8")

add_h2("Évolution mensuelle du chiffre d'affaires")
doc.add_picture(fr"{FIG}\01_ca_mensuel.png", width=Cm(16.5))

doc.add_page_break()

# ---------- PAGE 2 : CATEGORIE + VILLE ----------
add_h2("Analyse par catégorie — quelle catégorie prioriser ?")
add_body(
    "Électronique génère 74,6 % du CA total (1,87 M$) mais affiche la marge la plus faible (20 %) et le "
    "taux de retour le plus élevé (13,8 %, contre 2,8 % à 7,4 % pour les autres catégories). Beauté est à "
    "l'inverse la catégorie la plus rentable en marge (50 %) et la plus fiable (2,8 % de retour) mais reste "
    "marginale en volume (3 % du CA)."
)
doc.add_picture(fr"{FIG}\02_categorie.png", width=Cm(16.5))

add_h2("Analyse géographique — où investir davantage ?")
add_body(
    "Kinshasa domine largement (752 k$ de CA, 130 k$ de profit) et reste la priorité d'investissement. "
    "Douala combine la meilleure croissance du réseau (+76 %) et un taux d'annulation anormal de 12,9 % "
    "(quasi 0 % partout ailleurs) : un problème opérationnel local freine probablement un marché à fort "
    "potentiel. Libreville recule (-46 %) et doit être réévaluée avant tout nouvel investissement marketing. "
    "Brazzaville (+60 %) est un marché émergent à surveiller."
)
doc.add_picture(fr"{FIG}\03_ville.png", width=Cm(15.5))

doc.add_page_break()

# ---------- PAGE 3 : MARKETING + CLIENTS ----------
add_h2("Analyse marketing — quel canal mérite plus de budget ?")
add_body(
    "Email affiche un ROI de 225x pour seulement 2 349 $ investis — le plus petit budget des 4 canaux. "
    "À l'inverse, Instagram Ads consomme le plus gros budget (37 637 $) pour un ROI de seulement 24x, et "
    "Influenceur affiche le ROI le plus faible (21x). Google Ads est solide (ROI 49x). Un rééquilibrage "
    "progressif du budget vers Email et Google Ads, sans couper Instagram (premier contributeur de CA en "
    "volume), maximiserait le retour marginal du budget marketing."
)
doc.add_picture(fr"{FIG}\04_roi_canal.png", width=Cm(16.5))

add_h2("Analyse clients — comment améliorer la rétention ?")
add_body(
    "AfriMarket compte 1 747 clients, dont 74 % sont déjà récurrents (plus d'une commande) — un socle de "
    "fidélité solide. Cependant, les 350 clients les plus rentables (20 % de la base) concentrent 64,5 % du "
    "CA total : la valeur de l'entreprise repose sur une minorité de clients qu'il est prioritaire de "
    "protéger et développer via un programme de fidélisation dédié."
)
doc.add_picture(fr"{FIG}\05_pareto.png", width=Cm(15.5))

doc.add_page_break()

# ---------- PAGE 4 : RECOMMANDATIONS ----------
add_title("5 recommandations stratégiques", size=16, color=NAVY, space_after=8)

recos = [
    ("1. Sécuriser la marge sur Électronique",
     "Pilier du CA (74,6 %) mais marge la plus faible (20 %) et retours les plus fréquents (13,8 %). "
     "Négocier les coûts fournisseurs et traiter la cause des retours en priorité : effet de levier le "
     "plus élevé sur le profit total."),
    ("2. Réduire le taux de retour Électronique",
     "Cause principale du taux de retour global (8,1 %). Un audit qualité/SAV ciblé (fiabilité produit, "
     "clarté des fiches produit, réactivité du SAV) est le levier de rentabilité le plus direct identifié."),
    ("3. Réallouer le budget marketing vers Email et Google Ads",
     "Email : ROI 225x pour 2 349 $ investis (sous-financé). Instagram Ads : ROI 24x pour 37 637 $ investis "
     "(sur-financé). Transférer 15-20 % du budget Instagram/Influenceur vers Email/Google Ads, en test A/B "
     "avant généralisation."),
    ("4. Corriger l'anomalie opérationnelle de Douala et réévaluer Libreville",
     "Douala (+76 % de croissance) subit un taux d'annulation de 12,9 % à investiguer en urgence (paiement, "
     "stock, livraison). Libreville recule de 46 % et ne doit pas recevoir de budget supplémentaire tant "
     "que la cause du déclin n'est pas identifiée."),
    ("5. Lancer un programme de fidélisation VIP",
     "Les 350 clients Pareto (20 % de la base, 64,5 % du CA) doivent être protégés par un programme dédié "
     "(offres exclusives, service prioritaire) : le ROI de rétention y est supérieur à celui de l'acquisition."),
]

for title, detail in recos:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = BLUE
    add_body(detail)

doc.add_page_break()

# ---------- PAGE 5 : CONCLUSION ----------
add_title("Conclusion business orientée action", size=16, color=NAVY, space_after=8)
add_body(
    "Avec 2,51 M$ de CA et 430 k$ de profit net estimé sur 6 mois, AfriMarket a un modèle rentable mais "
    "avec trois fuites de rentabilité claires et quantifiées : un taux de retour Électronique presque double "
    "de la moyenne qui érode la marge de sa catégorie la plus vendue, un budget marketing mal réparti (Email "
    "sous-investi malgré un ROI 9x supérieur à Instagram Ads), et un problème opérationnel localisé à Douala "
    "qui freine la ville à plus forte croissance du réseau."
)

add_h2("Plan d'action à 30 jours")
add_bullet("Semaine 1 : audit qualité/SAV Électronique + audit opérationnel des annulations à Douala.")
add_bullet("Semaine 2 : transfert de 15-20 % du budget Instagram Ads/Influenceur vers Email/Google Ads, "
           "en test A/B sur un mois avant généralisation.")
add_bullet("Semaines 3-4 : lancement pilote du programme de fidélisation VIP sur les 350 clients Pareto identifiés.")

add_body(
    "Ces trois actions ciblent directement les 4 signaux remontés par la direction (variations de CA, taux "
    "de retour, dépenses marketing, écarts de performance géographique) avec un impact mesurable dès le mois "
    "suivant, sans investissement supplémentaire net : il s'agit de réallocation de ressources existantes, "
    "pas d'une augmentation de budget.",
    bold=True
)

doc.save(OUT)
print("Résumé exécutif sauvegardé :", OUT)
